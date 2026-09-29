from datetime import datetime
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock
from openpyxl import load_workbook
from backend.service.reports_service import ReportsService


async def test_monthly_report_uses_requested_month_and_warehouse():
    repo = AsyncMock()
    repo.get_operations_by_year_and_warehouse.return_value = (
        [
            SimpleNamespace(
                scheduled_at=datetime(2024, 2, 29),
                warehouse_id="a",
                supplier="Supplier",
            )
        ],
        [
            SimpleNamespace(
                scheduled_at=datetime(2024, 2, 29),
                warehouse_id="a",
                customer="Supplier",
            )
        ],
    )
    data = await ReportsService(repo).generate_monthly_report(2024, "a", [2])
    workbook = load_workbook(BytesIO(data))
    assert workbook.sheetnames == ["Февраль"]
    assert workbook.active.cell(1, 31).value == 29
    assert workbook.active.cell(2, 31).value == "ДТ+ОТ"
